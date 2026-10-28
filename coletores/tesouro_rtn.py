"""Tesouro Nacional, Resultado do Tesouro Nacional (RTN): série histórica mensal.

Dataset CKAN "resultado-do-tesouro-nacional" (id ab56485b-9c40-4efb-8563-9ce3e1973c4b).
O recurso XLSX muda de nome a cada edição (ex.: seriehistoricajul26.xlsx), por isso o endereço
é obtido no package_show a cada coleta.

Segundo o documento de metadados do Tesouro: coluna A com o código da linha, coluna B com a
descrição ("desdobramento") e, da coluna C em diante, valores mensais desde jan/1997,
em R$ milhões correntes.

Indicador gravado: Receita Total do Governo Central acumulada em 12 meses, em R$ bilhões.
O acumulado é a soma dos 12 valores mensais oficiais mais recentes (cálculo, não estimativa);
só é gravado quando os 12 meses existem. Substitui o "Impostômetro", que é uma estimativa da
Associação Comercial de São Paulo e não é fonte oficial.

Layout conferido em 28/10/2026 (seriehistoricajul26.xlsx): aba "1.1", cabeçalho na linha 5
("Discriminação" + meses como data), linha 6 = "1. RECEITA TOTAL 1/" (com nota de rodapé).
Se o rótulo ou o cabeçalho mudarem, LayoutErro e nada é gravado.
"""
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from urllib.parse import quote

from coletores.util import Desativado, LayoutErro, http_get, http_json, num

CKAN = "https://www.tesourotransparente.gov.br/ckan/api/3/action/package_show?id={}"
MESES = {m: i for i, m in enumerate(
    ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"], start=1)}


def _norm(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    return " ".join("".join(c for c in s if unicodedata.category(c) != "Mn").lower().split())


def mes_cabecalho(v):
    if isinstance(v, (datetime, date)):
        return date(v.year, v.month, 1)
    if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() and 30000 < v < 80000:
        d = date(1899, 12, 30) + timedelta(days=int(v))  # data como número de série do Excel
        return date(d.year, d.month, 1)
    s = _norm(v)
    m = re.fullmatch(r"([a-z]{3})[a-z]*[/\-\s]*(\d{2}|\d{4})", s)
    if m and m.group(1) in MESES:
        a = int(m.group(2))
        return date((a + 2000 if a < 70 else a + 1900) if a < 100 else a, MESES[m.group(1)], 1)
    return None


def parse_linhas(linhas, rotulo, url, divisor=1000):
    alvo = _norm(rotulo)
    cab = col_mes = None
    for i, l in enumerate(linhas[:40]):
        meses = {j: mes_cabecalho(c) for j, c in enumerate(l)}
        meses = {j: m for j, m in meses.items() if m}
        if len(meses) >= 24:
            cab, col_mes = i, meses
            break
    if cab is None:
        raise LayoutErro("RTN: linha de cabeçalho com os meses não encontrada")
    linha = None
    for l in linhas[cab + 1:]:
        textos = [_norm(c) for c in l[:3] if isinstance(c, str)]
        # "1. RECEITA TOTAL 1/": remove o código à esquerda e as notas de rodapé ("1/") à direita
        if any(re.sub(r"(\s+\d+/)+$", "", re.sub(r"^[\d.\s\-]+", "", t)) == alvo for t in textos):
            linha = l
            break
    if linha is None:
        raise LayoutErro(f"RTN: linha '{rotulo}' não encontrada")
    mensal = {}
    for j, m in col_mes.items():
        v = num(linha[j]) if j < len(linha) else None
        if v is not None:
            mensal[m] = v
    if len(mensal) < 24:
        raise LayoutErro("RTN: menos de 24 meses com valor na linha")
    ordem = sorted(mensal)
    obs = []
    for k in range(11, len(ordem)):
        janela = ordem[k - 11:k + 1]
        if (janela[-1].year - janela[0].year) * 12 + janela[-1].month - janela[0].month != 11:
            continue  # meses faltando: não acumula
        obs.append({"ref": janela[-1].isoformat(), "valor": sum(mensal[m] for m in janela) / divisor, "url": url})
    return obs


def linhas_planilha(conteudo, aba):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    nomes = {_norm(n): n for n in wb.sheetnames}
    nome = nomes.get(_norm(aba)) or next((n for k, n in nomes.items() if k.startswith(_norm(aba))), None)
    if not nome:
        raise LayoutErro(f"RTN: aba '{aba}' não encontrada; abas: {wb.sheetnames[:15]}")
    return list(wb[nome].iter_rows(values_only=True))


def coletar(ind, primeira, hoje, _json=http_json, _get=http_get):
    p = ind["params"]
    if not p.get("dataset"):
        raise Desativado("RTN: dataset não configurado")
    recursos = (_json(CKAN.format(quote(p["dataset"]))).get("result") or {}).get("resources") or []
    xls = [r for r in recursos if str(r.get("format", "")).upper() == "XLSX" and "serie" in _norm(r.get("name"))]
    if not xls:
        raise LayoutErro("RTN: recurso XLSX da série histórica não encontrado no dataset")
    url = xls[0]["url"].replace("http://", "https://")
    obs = parse_linhas(linhas_planilha(_get(url, timeout=120), p["aba"]), p["rotulo"], url, p.get("divisor", 1000))
    return obs[-24:] if primeira else obs[-3:]
