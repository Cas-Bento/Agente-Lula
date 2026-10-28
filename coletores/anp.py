"""ANP, Levantamento de Preços de Combustíveis (LPC): série histórica semanal, Brasil.

Arquivo oficial (atualizado semanalmente pela ANP):
  .../shlp/semanal/semanal-brasil-desde-2013.xlsx
Layout conferido em 28/10/2026: aba única "BRASIL - DESDE 30.12.2012"; cabeçalho na linha
que contém "DATA INICIAL | DATA FINAL | PRODUTO | NÚMERO DE POSTOS PESQUISADOS |
UNIDADE DE MEDIDA | PREÇO MÉDIO REVENDA | ...". Produtos: "GASOLINA COMUM",
"ETANOL HIDRATADO", "OLEO DIESEL S10", "OLEO DIESEL" (S500) etc.

Indicador gravado: PREÇO MÉDIO REVENDA (R$/l) do produto configurado; ref = DATA FINAL da semana.
"""
import io
import re
import unicodedata
from datetime import date, datetime, timedelta

from coletores.util import LayoutErro, http_get, num

URL = ("https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/"
       "precos-revenda-e-de-distribuicao-combustiveis/shlp/semanal/semanal-brasil-desde-2013.xlsx")
_cache = {}  # o mesmo arquivo atende gasolina, etanol e diesel: baixa uma vez por execução


def _norm(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    return " ".join("".join(c for c in s if unicodedata.category(c) != "Mn").upper().split())


def _data(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 30000 < v < 80000:  # número de série do Excel
        return date(1899, 12, 30) + timedelta(days=int(v))
    m = re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})", str(v or "").strip())
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


def linhas_planilha(conteudo):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    return list(wb.worksheets[0].iter_rows(values_only=True))


def parse_linhas(linhas, produto, url, meses=None):
    cab = None
    for i, l in enumerate(linhas):
        n = [_norm(c) for c in l]
        if "PRODUTO" in n and "DATA FINAL" in n:
            cab = i
            idx = {k: n.index(k) for k in ("DATA FINAL", "PRODUTO")}
            col_p = next((j for j, c in enumerate(n) if c == "PRECO MEDIO REVENDA"), None)
            col_u = next((j for j, c in enumerate(n) if c == "UNIDADE DE MEDIDA"), None)
            break
    if cab is None or col_p is None:
        raise LayoutErro("ANP: cabeçalho com DATA FINAL / PRODUTO / PREÇO MÉDIO REVENDA não encontrado")
    alvo = _norm(produto)
    obs = {}
    for l in linhas[cab + 1:]:
        if len(l) <= max(col_p, idx["PRODUTO"], idx["DATA FINAL"]) or _norm(l[idx["PRODUTO"]]) != alvo:
            continue
        d = _data(l[idx["DATA FINAL"]])
        v = num(l[col_p])
        if d is None or v is None:
            continue  # semana sem pesquisa ("-"): não completar
        if col_u is not None and _norm(l[col_u]) not in ("R$/L", ""):
            raise LayoutErro(f"ANP: unidade inesperada para {produto}: {l[col_u]!r}")
        obs[d.isoformat()] = {"ref": d.isoformat(), "valor": v, "url": url}
    if not obs:
        raise LayoutErro(f"ANP: produto '{produto}' sem linhas na planilha")
    saida = sorted(obs.values(), key=lambda o: o["ref"])
    if meses:
        corte = (date.fromisoformat(saida[-1]["ref"]) - timedelta(days=31 * meses)).isoformat()
        saida = [o for o in saida if o["ref"] >= corte]
    return saida


def coletar(ind, primeira, hoje, _get=http_get):
    if URL not in _cache:
        _cache[URL] = linhas_planilha(_get(URL, timeout=90))
    return parse_linhas(_cache[URL], ind["params"]["produto"], URL, meses=24 if primeira else 2)
